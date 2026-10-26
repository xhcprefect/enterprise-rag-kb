"""真实临时文件验证结构位置、内容保留、尺寸和表格顺序，无模型/数据库。"""
import tempfile
import unittest
from pathlib import Path

from app.services.structured_documents import iter_structured_chunks


class StructuredDocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def chunks(self, path, size=120, overlap=20):
        chunks = list(iter_structured_chunks(path, size, overlap))
        self.assertTrue(chunks)
        self.assertTrue(all(0 < len(chunk.content) <= size for chunk in chunks))
        self.assertTrue(all(0 <= chunk.progress <= 1 for chunk in chunks))
        self.assertEqual([chunk.progress for chunk in chunks], sorted(chunk.progress for chunk in chunks))
        return chunks

    def test_long_single_line_is_bounded_and_preserves_every_character(self):
        text = "".join(chr(0x5000 + index) for index in range(2000))
        path = self.root / "long.txt"
        path.write_text(text, encoding="utf-8")
        chunks = self.chunks(path, 64, 12)
        self.assertGreater(len(chunks), 20)
        self.assertTrue(set(text).issubset(set("".join(chunk.content for chunk in chunks))))
        self.assertTrue(all(chunk.metadata["line_start"] == chunk.metadata["line_end"] == 1 for chunk in chunks))

    def test_headings_and_articles_are_not_merged_across_sections(self):
        path = self.root / "law.md"
        path.write_text("# 第一章\n第一条 " + "审批要求。" * 50 + "\n# 第二章\n第二条 保存记录。\n", encoding="utf-8")
        chunks = self.chunks(path)
        first = [chunk for chunk in chunks if "第一条" in chunk.metadata.get("section", "")]
        second = [chunk for chunk in chunks if "第二条" in chunk.metadata.get("section", "")]
        self.assertTrue(first and second)
        self.assertTrue(all("第一章" in chunk.content for chunk in first))
        self.assertTrue(all("第二章" not in chunk.content for chunk in first))
        self.assertTrue(all("第一章" not in chunk.content for chunk in second))

    def test_csv_repeats_headers_keeps_empty_cells_and_records_rows(self):
        path = self.root / "data.csv"
        rows = [f"id{index:03},,city{index:03}" for index in range(15)]
        path.write_text("编号,年龄,城市\n" + "\n".join(rows), encoding="utf-8")
        chunks = self.chunks(path, 120, 15)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertIn("编号", chunk.content)
            self.assertIn("年龄", chunk.content)
            self.assertGreaterEqual(chunk.metadata["row_start"], 2)
        joined = "".join(chunk.content for chunk in chunks)
        self.assertIn("id000 |  | city000", joined)
        self.assertIn("id014", joined)

    def test_markdown_tables_repeat_headers_and_keep_original_line_ranges(self):
        path = self.root / "table.md"
        path.write_text("# 费用\n| 项目 | 金额 |\n| --- | ---: |\n" +
                        "\n".join(f"| item{index:02} | {index}元 |" for index in range(20)) +
                        "\n\n表格之后的说明。\n", encoding="utf-8")
        chunks = self.chunks(path, 120, 12)
        tables = [chunk for chunk in chunks if chunk.metadata.get("table_index") == 1]
        self.assertGreater(len(tables), 2)
        self.assertTrue(all("| 项目 | 金额 |" in chunk.content for chunk in tables))
        self.assertTrue(all(chunk.metadata["line_start"] >= 4 and chunk.metadata["line_end"] <= 23 for chunk in tables))
        self.assertIn("item19", "".join(chunk.content for chunk in tables))
        self.assertTrue(any("表格之后" in chunk.content and not chunk.metadata.get("table_index") for chunk in chunks))

    def test_header_only_markdown_table_and_code_fence_headings_are_preserved(self):
        path = self.root / "code.md"
        path.write_text("# 正式章节\n```md\n# 不应成为章节\n| a | b |\n| --- | --- |\n```\n正文\n", encoding="utf-8")
        chunks = self.chunks(path)
        self.assertTrue(all(chunk.metadata.get("section") == "正式章节" for chunk in chunks))
        self.assertFalse(any(chunk.metadata.get("table_index") for chunk in chunks))
        table = self.root / "header-only.md"
        table.write_text("| 名称 | 金额 |\n| --- | --- |\n", encoding="utf-8")
        self.assertIn("名称", "".join(chunk.content for chunk in self.chunks(table)))

    def test_excel_records_sheet_and_original_row_numbers(self):
        from openpyxl import Workbook
        path = self.root / "sheets.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.title = "预算"
        first.append(["名称", "金额"])
        first.append(["研发", 100])
        other = workbook.create_sheet("人数")
        other.append(["部门", "人数"])
        other.append(["销售", 3])
        workbook.save(path)
        workbook.close()
        chunks = self.chunks(path)
        self.assertEqual({chunk.metadata["sheet"] for chunk in chunks}, {"预算", "人数"})
        self.assertTrue(all(chunk.metadata["row_start"] == 2 for chunk in chunks))

    def test_docx_preserves_interleaved_paragraphs_and_table(self):
        from docx import Document
        path = self.root / "ordered.docx"
        document = Document()
        document.add_heading("审批流程", 1)
        document.add_paragraph("表格之前")
        table = document.add_table(rows=2, cols=3)
        for cell, text in zip(table.rows[0].cells, ["A", "B", "C"]):
            cell.text = text
        table.cell(1, 0).text, table.cell(1, 2).text = "left", "right"
        document.add_paragraph("表格之后")
        document.save(path)
        chunks = self.chunks(path)
        joined = "\n".join(chunk.content for chunk in chunks)
        self.assertLess(joined.index("表格之前"), joined.index("left |  | right"))
        self.assertLess(joined.index("left |  | right"), joined.index("表格之后"))
        self.assertTrue(any(chunk.metadata.get("table_index") == 1 for chunk in chunks))

    def test_pdf_retains_actual_page_numbers_without_crossing_pages(self):
        import pymupdf
        path = self.root / "pages.pdf"
        document = pymupdf.open()
        for text in ("first page evidence", "second page evidence"):
            document.new_page().insert_text((50, 50), text)
        document.save(path)
        document.close()
        chunks = self.chunks(path)
        self.assertEqual({chunk.metadata["page_start"] for chunk in chunks}, {1, 2})
        self.assertTrue(all(chunk.metadata["page_start"] == chunk.metadata["page_end"] for chunk in chunks))

    def test_empty_nul_and_unsupported_files_are_handled_explicitly(self):
        for name, text in [("empty.txt", " \n\x00"), ("bad.bin", "content")]:
            path = self.root / name
            path.write_text(text, encoding="utf-8")
            with self.assertRaises(ValueError):
                list(iter_structured_chunks(path, 80, 10))
        path = self.root / "nul.txt"
        path.write_text("a\x00b", encoding="utf-8")
        self.assertEqual(self.chunks(path)[0].content, "ab")

    def test_very_long_header_cannot_exceed_chunk_budget(self):
        path = self.root / "header.csv"
        path.write_text("超长表头" * 100 + ",B\nvalue,other\n", encoding="utf-8")
        chunks = self.chunks(path, 40, 10)
        self.assertTrue(any(chunk.metadata.get("prefix_truncated") for chunk in chunks))
        self.assertIn("value", "".join(chunk.content for chunk in chunks))
