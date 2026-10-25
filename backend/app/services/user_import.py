"""批量导入用户：解析表格 -> 逐行校验 -> 生成随机密码 -> 落库 -> 生成回执。

依赖只有 openpyxl 与标准库。openpyxl 本来就在 requirements 里（文档解析用），
所以这个功能不需要新增任何依赖；CSV 用标准库 ``csv`` 而不是 pandas，
是因为导入需要的是「单元格原样字符串」——pandas 的类型推断会把纯数字用户名
变成 int、把空单元格变成 NaN，反而要来回转换。

设计取舍：

* **尽力而为，不是全有或全无**。一行不合法不影响其它行，每行都给出结果；
  但落库用的是逐行 SAVEPOINT，所以两处并发导入同名用户也只会有一个人失败，
  不会把整批回滚掉。
* **密码只在本次响应里出现一次**。库里存的是 bcrypt 哈希，事后无法还原，
  所以回执（xlsx）与接口响应必须同时带明文，前端展示一次后由管理员负责分发。
* **用户名匹配与单条创建接口完全一致**（精确匹配、区分大小写），
  避免「手动能建、导入说重复」这种自相矛盾的行为。
"""
import base64
import csv
import io
import logging
import os
import re
import secrets
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.models import User
from app.services.parser import detect_encoding

logger = logging.getLogger(__name__)

# 只收这两种：.xls 是旧版二进制格式，openpyxl 读不了，需要额外装 xlrd
SUPPORTED_IMPORT_EXTENSIONS = {".xlsx", ".csv"}

# 单次导入上限。一是防止一个手滑把几万行塞进来（bcrypt 每行 200ms 左右），
# 二是让请求耗时有个可预期的上界。
MAX_IMPORT_ROWS = 500
# 用户名单是纯文本，5MB 足够 500 行的几十倍富余
MAX_IMPORT_BYTES = 5 * 1024 * 1024

# 与 schemas/user.py 的 UserCreate 保持一致；不一致会让「导入能过、手动建不了」
USERNAME_PATTERN = re.compile(r"^[A-Za-z0-9_.-]+$")
USERNAME_MIN, USERNAME_MAX = 2, 64
DISPLAY_NAME_MAX = 128
EMAIL_MAX = 255

# ---------- 随机密码 ----------

# 去掉了容易看错的 0 O 1 l I；不用符号是因为密码要靠人工抄写/口述分发，
# 键盘布局差异（尤其是符号键位置）比多几位长度更容易造成登录失败。
# 54 个字符 × 14 位 ≈ 80 bit 熵，远超需要，同时好读好打。
_PWD_LOWER = "abcdefghjkmnpqrstuvwxyz"
_PWD_UPPER = "ABCDEFGHJKMNPQRSTUVWXYZ"
_PWD_DIGITS = "23456789"
_PWD_ALPHABET = _PWD_LOWER + _PWD_UPPER + _PWD_DIGITS
PASSWORD_LENGTH = 14


def generate_password(length: int = PASSWORD_LENGTH) -> str:
    """生成随机初始密码，保证大写/小写/数字三类都至少出现一次。"""
    while True:
        pwd = "".join(secrets.choice(_PWD_ALPHABET) for _ in range(length))
        if (
            any(c in _PWD_LOWER for c in pwd)
            and any(c in _PWD_UPPER for c in pwd)
            and any(c in _PWD_DIGITS for c in pwd)
        ):
            return pwd


def _hash_many(passwords: list[str]) -> list[str]:
    """并发做 bcrypt。

    bcrypt 是纯 C 调用且会释放 GIL，实测 20 条串行 4.26s、4 线程 1.08s。
    批量导入是一次性的管理操作，临时借几个线程换掉几十秒等待是值得的；
    上限卡在 4，避免和问答用的线程池抢 CPU（服务器只有 2 核，实际就是 2）。
    """
    workers = max(1, min(4, os.cpu_count() or 1))
    if len(passwords) < 8 or workers == 1:
        return [hash_password(p) for p in passwords]
    with ThreadPoolExecutor(max_workers=workers) as pool:
        return list(pool.map(hash_password, passwords))


# ---------- 表头识别 ----------

# 表头做「去空格/下划线/横线 + 转小写」归一化后比对，所以这里也按归一化形式写。
_HEADER_ALIASES: dict[str, set[str]] = {
    "username": {
        "用户名", "账号", "帐号", "登录名", "登录账号", "登陆名", "用户",
        "username", "user", "account", "login", "loginname",
    },
    "display_name": {
        "姓名", "显示名称", "名称", "昵称", "真实姓名", "用户姓名",
        "displayname", "name", "fullname", "realname", "nickname",
    },
    "email": {"邮箱", "电子邮件", "邮件", "邮箱地址", "email", "mail", "e-mail"},
    "role": {"角色", "权限", "角色权限", "用户角色", "role", "roles", "type"},
}

_ROLE_ALIASES: dict[str, set[str]] = {
    "admin": {"管理员", "管理", "超管", "超级管理员", "admin", "administrator", "superuser"},
    "user": {
        "普通用户", "普通", "用户", "成员", "员工",
        "user", "normal", "member", "staff", "employee",
    },
}


def _norm_header(value: str) -> str:
    return value.strip().lower().replace(" ", "").replace("_", "").replace("-", "")


def _norm_role(value: str) -> str | None:
    """角色归一化；空值按普通用户处理，无法识别返回 None。"""
    text = _norm_header(value)
    if not text:
        return "user"
    for role, aliases in _ROLE_ALIASES.items():
        if text in aliases:
            return role
    return None


# ---------- 结果结构 ----------


@dataclass
class CreatedUser:
    row: int
    username: str
    display_name: str
    email: str | None
    role: str
    password: str


@dataclass
class ImportFailure:
    row: int
    username: str
    reason: str


@dataclass
class ImportOutcome:
    total: int
    created: list[CreatedUser] = field(default_factory=list)
    failed: list[ImportFailure] = field(default_factory=list)
    receipt_filename: str = ""
    receipt_bytes: bytes = b""


class ImportFormatError(ValueError):
    """表格本身读不了（格式、表头、行数超限）——对应的是一次 400。"""


# ---------- 读表格 ----------


def _cell_text(value: object) -> str:
    """单元格转字符串：保留原样，不做任何类型推断。"""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        # Excel 里纯数字单元格读出来是 float，1001.0 要还原成 "1001"
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def _read_csv_rows(data: bytes) -> list[list[str]]:
    encoding = detect_encoding(data[:64 * 1024])
    try:
        text = data.decode(encoding)
    except UnicodeDecodeError as exc:
        raise ImportFormatError(
            f"CSV 编码无法识别（尝试了 {encoding}），请另存为 UTF-8 或 GBK 后重试"
        ) from exc
    # Excel 导出的「UTF-8 CSV」首字节常带 BOM。utf-8 解码器会把它解成 U+FEFF，
    # 粘在第一个表头上导致识别不到「用户名」列，所以这里先剥掉。
    text = text.lstrip("﻿")

    rows = list(csv.reader(io.StringIO(text)))
    # 部分 Excel 在中文/欧洲区域设置下导出的是分号分隔；只在「看起来确实不是逗号」
    # 时才换分隔符重读，避免把正常的逗号文件读坏。
    if rows and all(len(r) == 1 for r in rows if r):
        if ";" in text.splitlines()[0]:
            rows = list(csv.reader(io.StringIO(text), delimiter=";"))
        elif "\t" in text.splitlines()[0]:
            rows = list(csv.reader(io.StringIO(text), delimiter="\t"))
    return [[_cell_text(c) for c in row] for row in rows]


def _read_xlsx_rows(data: bytes) -> list[list[str]]:
    try:
        # read_only：不把整张表建进内存；data_only：取公式的缓存值而不是公式本身
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl 抛的异常类型很杂
        raise ImportFormatError(f"Excel 文件无法解析：{exc}") from exc
    try:
        ws = wb.active
        if ws is None:
            raise ImportFormatError("Excel 里没有工作表")
        return [[_cell_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_rows(data: bytes, ext: str) -> list[list[str]]:
    return _read_csv_rows(data) if ext == ".csv" else _read_xlsx_rows(data)


def _resolve_columns(header: list[str]) -> dict[str, int]:
    """把表头映射成字段名 -> 列下标。未知列忽略，用户名列必须有。"""
    columns: dict[str, int] = {}
    for index, cell in enumerate(header):
        key = _norm_header(cell)
        if not key:
            continue
        for field_name, aliases in _HEADER_ALIASES.items():
            if field_name in columns:
                continue
            if key in aliases:
                columns[field_name] = index
                break
    if "username" not in columns:
        raise ImportFormatError(
            "表头里找不到「用户名」列。可用的列名：用户名/账号/登录名（username），"
            "姓名（display_name），邮箱（email），角色（role）"
        )
    return columns


def _cell(row: list[str], columns: dict[str, int], name: str) -> str:
    index = columns.get(name)
    if index is None or index >= len(row):
        return ""
    return row[index]


@dataclass
class _Pending:
    """通过校验、准备落库的一行。"""
    row: int
    username: str
    display_name: str
    email: str | None
    role: str


def _parse(data: bytes, ext: str) -> tuple[list[_Pending], list[ImportFailure]]:
    """读表 + 逐行校验，返回（可落库的行，校验失败的行）。"""
    rows = _read_rows(data, ext)
    if not rows:
        raise ImportFormatError("文件是空的")

    columns = _resolve_columns(rows[0])

    # 先去掉整行空白（Excel 常在末尾留一堆空行）。注意要把原始行号一起带上：
    # 直接 enumerate 过滤后的列表会让后面的行号整体前移，回执里的「表格行号」
    # 就对不上用户在 Excel 里看到的行号，等于白给。
    numbered = [
        (index, row) for index, row in enumerate(rows[1:], start=2) if any(row)
    ]  # 第 1 行是表头，数据从第 2 行起
    if not numbered:
        raise ImportFormatError("表头下面没有任何数据行")
    if len(numbered) > MAX_IMPORT_ROWS:
        raise ImportFormatError(
            f"本次共 {len(numbered)} 行数据，超过单次导入上限 {MAX_IMPORT_ROWS} 行，请拆分后重试"
        )

    pending: list[_Pending] = []
    failures: list[ImportFailure] = []
    seen: dict[str, int] = {}

    for offset, row in numbered:
        username = _cell(row, columns, "username")
        display = _cell(row, columns, "display_name")
        email = _cell(row, columns, "email")
        role_raw = _cell(row, columns, "role")

        if not username:
            failures.append(ImportFailure(offset, "", "用户名为空"))
            continue
        if not (USERNAME_MIN <= len(username) <= USERNAME_MAX):
            failures.append(
                ImportFailure(offset, username, f"用户名长度需在 {USERNAME_MIN}-{USERNAME_MAX} 之间")
            )
            continue
        if not USERNAME_PATTERN.match(username):
            failures.append(
                ImportFailure(offset, username, "用户名只能包含字母、数字、下划线、点、横线")
            )
            continue
        if username in seen:
            failures.append(
                ImportFailure(offset, username, f"与第 {seen[username]} 行的用户名重复")
            )
            continue
        if len(display) > DISPLAY_NAME_MAX:
            failures.append(
                ImportFailure(offset, username, f"姓名超过 {DISPLAY_NAME_MAX} 个字符")
            )
            continue
        if email and len(email) > EMAIL_MAX:
            failures.append(
                ImportFailure(offset, username, f"邮箱超过 {EMAIL_MAX} 个字符")
            )
            continue
        role = _norm_role(role_raw)
        if role is None:
            failures.append(
                ImportFailure(offset, username, f"角色「{role_raw}」无法识别，可填：管理员 / 普通用户")
            )
            continue

        seen[username] = offset
        pending.append(
            _Pending(
                row=offset,
                username=username,
                display_name=display or username,
                email=email or None,
                role=role,
            )
        )

    return pending, failures


# ---------- 落库 ----------


def _drop_existing(db: Session, pending: list[_Pending], failures: list[ImportFailure]) -> list[_Pending]:
    """一次性查重，把已存在的用户名挪进失败列表。"""
    if not pending:
        return []
    existing = set(
        db.scalars(
            select(User.username).where(User.username.in_([p.username for p in pending]))
        ).all()
    )
    if not existing:
        return pending

    kept: list[_Pending] = []
    for item in pending:
        if item.username in existing:
            failures.append(ImportFailure(item.row, item.username, "用户名已存在，已跳过"))
        else:
            kept.append(item)
    return kept


# ---------- 回执 ----------

_HEADER_FILL = PatternFill("solid", fgColor="EEF2FF")
_HEADER_FONT = Font(bold=True)
_TIP_FONT = Font(color="9A3412", italic=True)


def _write_header(ws, titles: list[str]) -> None:
    ws.append(titles)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _autosize(ws, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def build_receipt(created: list[CreatedUser], failed: list[ImportFailure], source_name: str) -> tuple[str, bytes]:
    """生成导入回执：成功者带初始密码，失败者带原因，各一张表。"""
    wb = Workbook()

    ws = wb.active
    ws.title = "导入成功"
    _write_header(ws, ["用户名", "姓名", "邮箱", "角色", "初始密码"])
    for item in created:
        ws.append([item.username, item.display_name, item.email or "",
                   "管理员" if item.role == "admin" else "普通用户", item.password])
    # 密码列必须是文本格式，否则 Excel 会把 2A3b... 之类按数字/科学计数法处理
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "@"
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "@"
    if created:
        tip_row = ws.max_row + 2
        ws.cell(row=tip_row, column=1, value=(
            "以上为初始密码，仅本次可见（数据库中只存哈希，无法再次导出）。"
            "请通过可靠渠道分发给本人，并提醒登录后立即修改。"
        )).font = _TIP_FONT
    _autosize(ws, {"A": 22, "B": 16, "C": 30, "D": 10, "E": 20})

    ws2 = wb.create_sheet("导入失败")
    _write_header(ws2, ["表格行号", "用户名", "失败原因"])
    for item in failed:
        ws2.append([item.row, item.username, item.reason])
    _autosize(ws2, {"A": 10, "B": 22, "C": 48})

    ws3 = wb.create_sheet("说明")
    ws3.append(["源文件", source_name or "（未命名）"])
    ws3.append(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws3.append(["成功", len(created)])
    ws3.append(["失败", len(failed)])
    ws3.append([])
    ws3.append(["提示", "失败的行没有创建账号，修正后重新导入即可；已存在的用户名会被跳过。"])
    for cell in ws3["A"]:
        cell.font = _HEADER_FONT if cell.row <= 4 else Font()
    _autosize(ws3, {"A": 12, "B": 70})
    ws3.column_dimensions["B"].alignment = Alignment(wrap_text=True, vertical="top")

    wb.properties.title = "用户导入回执"
    wb.properties.creator = "Enterprise RAG Knowledge Base"

    # 回执没有业务公式。强制所有文本单元格为字符串，阻止 = 开头的外部数据变成公式。
    for sheet in wb:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    buffer = io.BytesIO()
    wb.save(buffer)
    filename = f"用户导入回执_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return filename, buffer.getvalue()


def build_template() -> tuple[str, bytes]:
    """导入模板：只有表头 + 一个下拉框，不放示例数据，免得被当成真实账号导进去。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "用户"
    _write_header(ws, ["用户名", "姓名", "邮箱", "角色"])

    # 角色列做下拉，省得填错还要回来改
    validation = DataValidation(
        type="list", formula1='"普通用户,管理员"', allow_blank=True,
        showDropDown=False,
    )
    validation.error = "只能选：普通用户 / 管理员"
    validation.errorTitle = "角色取值无效"
    validation.prompt = "留空默认为普通用户"
    validation.promptTitle = "角色"
    ws.add_data_validation(validation)
    validation.add(f"D2:D{MAX_IMPORT_ROWS + 1}")

    _autosize(ws, {"A": 22, "B": 16, "C": 30, "D": 12})

    tips = wb.create_sheet("填写说明")
    lines = [
        ("必填", "用户名（2-64 位，只能包含字母、数字、下划线、点、横线；登录时区分大小写）"),
        ("选填", "姓名（留空则与用户名相同）"),
        ("选填", "邮箱（留空即可）"),
        ("选填", "角色（普通用户 / 管理员，留空为普通用户）"),
        ("密码", "由系统为每个人随机生成，导入完成后下载回执查看，每人一个不同密码"),
        ("重复", "用户名已存在的行会被跳过并计入失败，不会覆盖已有账号"),
        ("格式", "支持 .xlsx 与 .csv；.xls（旧版 Excel）请先另存为 .xlsx"),
        ("上限", f"单次最多 {MAX_IMPORT_ROWS} 行"),
    ]
    tips.append(["项目", "说明"])
    for cell in tips[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for key, value in lines:
        tips.append([key, value])
    _autosize(tips, {"A": 10, "B": 76})
    for row in tips.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    wb.save(buffer)
    return "用户导入模板.xlsx", buffer.getvalue()


# ---------- 入口 ----------


def import_users(db: Session, data: bytes, filename: str) -> ImportOutcome:
    """解析并导入。数据库里的写入在这里提交；调用方只负责审计与响应。"""
    if not data:
        raise ImportFormatError("文件是空的")
    if len(data) > MAX_IMPORT_BYTES:
        raise ImportFormatError(
            f"文件超过 {MAX_IMPORT_BYTES // 1024 // 1024}MB 限制，请精简后重试"
        )

    ext = os.path.splitext(filename or "")[1].lower()
    if ext == ".xls":
        raise ImportFormatError("不支持旧版 .xls，请在 Excel 里另存为 .xlsx 或 .csv 后重试")
    if ext not in SUPPORTED_IMPORT_EXTENSIONS:
        raise ImportFormatError(
            f"不支持的文件类型 {ext or '（无扩展名）'}，仅支持 .xlsx 和 .csv"
        )

    pending, failures = _parse(data, ext)
    # 查重会往 failures 里追加，所以 total 要在查重之后再算
    pending = _drop_existing(db, pending, failures)
    total = len(pending) + len(failures)

    created: list[CreatedUser] = []
    if pending:
        passwords = [generate_password() for _ in pending]
        hashes = _hash_many(passwords)
        for item, pwd, pwd_hash in zip(pending, passwords, hashes):
            try:
                with db.begin_nested():
                    db.add(
                        User(
                            username=item.username,
                            display_name=item.display_name,
                            email=item.email,
                            password_hash=pwd_hash,
                            role=item.role,
                        )
                    )
                    db.flush()
            except IntegrityError:
                failures.append(ImportFailure(item.row, item.username, "用户名已存在，已跳过"))
                continue
            created.append(
                CreatedUser(
                    row=item.row,
                    username=item.username,
                    display_name=item.display_name,
                    email=item.email,
                    role=item.role,
                    password=pwd,
                )
            )
        db.commit()

    failures.sort(key=lambda f: f.row)

    receipt_name, receipt_bytes = build_receipt(created, failures, filename)
    logger.info(
        "批量导入用户: 文件=%s 共%d行 成功%d 失败%d",
        filename or "-", total, len(created), len(failures),
    )
    return ImportOutcome(
        total=total,
        created=created,
        failed=failures,
        receipt_filename=receipt_name,
        receipt_bytes=receipt_bytes,
    )


def receipt_b64(outcome: ImportOutcome) -> str:
    return base64.b64encode(outcome.receipt_bytes).decode("ascii")
