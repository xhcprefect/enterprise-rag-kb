"""有界结构化解析与分块，保留原文位置，表格和章节前缀在分块中重复。"""
import csv
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from app.services.parser import detect_encoding
from app.services.source_documents import merge_locations


@dataclass(frozen=True)
class StructuredChunk:
    content: str
    metadata: dict
    progress: float


@dataclass(frozen=True)
class _Block:
    text: str
    metadata: dict
    progress: float
    prefix: str = ""


class _Sections:
    def __init__(self):
        self.values = {}
        self.serial = 0

    def update(self, text, level=None):
        stripped = text.strip()
        markdown = re.match(r"^(#{1,6})\s+(.+)", stripped)
        legal = re.match(r"^(第[零一二三四五六七八九十百千万两0-9]+([编章节条款]))", stripped)
        if markdown:
            level, title = len(markdown.group(1)), markdown.group(2).rstrip("# ")
        elif legal:
            level = {"编": 1, "章": 2, "节": 3, "条": 4, "款": 5}[legal.group(2)]
            title = stripped[:160]
        elif level:
            title = stripped[:160]
        else:
            return
        self.values = {key: value for key, value in self.values.items() if key < level}
        self.values[level] = title
        self.serial += 1

    def metadata(self):
        if not self.values:
            return {}
        path = [value for _, value in sorted(self.values.items())]
        return {"section": " / ".join(path)[:512], "section_path": path,
                "section_key": f"s{self.serial}"}

    def prefix(self):
        return "\n".join(value for _, value in sorted(self.values.items()))


def _text_lines(path):
    with path.open("rb") as stream:
        encoding = detect_encoding(stream.read(65536))
    if encoding == "utf-8":
        encoding = "utf-8-sig"
    line, size = 1, max(1, path.stat().st_size)
    with path.open("r", encoding=encoding, errors="replace", newline="") as stream:
        while text := stream.readline(65536):
            text = text.replace("\r\n", "\n").replace("\r", "\n")
            end = line + text.count("\n") - (1 if text.endswith("\n") else 0)
            yield text, {"line_start": line, "line_end": max(line, end)}, min(1.0, stream.buffer.tell() / size)
            line += text.count("\n")


def _markdown_separator(text):
    cells = text.strip().strip("|").split("|")
    return "|" in text and all(re.fullmatch(r"\s*:?-{3,}:?\s*", cell) for cell in cells)


def _text_blocks(path):
    sections, lines, pending = _Sections(), iter(_text_lines(path)), None
    markdown, fence, table, table_no = path.suffix.lower() == ".md", None, None, 0
    while True:
        item = pending if pending is not None else next(lines, None)
        pending = None
        if item is None:
            if table is not None and not table["produced"]:
                yield _Block(table["header"], table["metadata"], table["progress"], sections.prefix())
            break
        text, location, progress = item
        if table is not None:
            if text.strip() and "|" in text:
                table["produced"] = True
                yield _Block(text, {**sections.metadata(), **location, "table_index": table_no}, progress,
                             "\n".join(part for part in (sections.prefix(), table["header"].rstrip()) if part))
                continue
            if not table["produced"]:
                yield _Block(table["header"], table["metadata"], table["progress"], sections.prefix())
            table = None
        marker = re.match(r"^[ \t]{0,3}(`{3,}|~{3,})", text) if markdown else None
        inside_fence = fence is not None
        if marker:
            value = marker.group(1)
            if fence is None:
                fence = value
            elif value[0] == fence[0] and len(value) >= len(fence) and not text[marker.end():].strip():
                fence = None
        if not inside_fence and not marker:
            sections.update(text)
            if markdown and "|" in text and not re.match(r"^\s*#{1,6}\s", text):
                following = next(lines, None)
                if following is not None and _markdown_separator(following[0]):
                    table_no += 1
                    table = {"header": text.rstrip("\n") + "\n" + following[0], "produced": False,
                             "metadata": {**sections.metadata(), **location,
                                          "line_end": following[1]["line_end"], "table_index": table_no},
                             "progress": following[2]}
                    continue
                pending = following
        yield _Block(text, {**sections.metadata(), **location}, progress, sections.prefix())


def _pdf_blocks(path):
    import pymupdf
    sections = _Sections()
    with pymupdf.open(path) as document:
        total = max(1, document.page_count)
        for page_no, page in enumerate(document, 1):
            for line in page.get_text("text").splitlines(keepends=True):
                sections.update(line)
                yield _Block(line, {**sections.metadata(), "page_start": page_no, "page_end": page_no},
                             page_no / total, sections.prefix())


def _cell(value):
    return "" if value is None else str(value).replace("\x00", "").replace("\n", " / ").replace("|", "\\|")


def _table_blocks(rows, base, progress, prefix=""):
    header = None
    produced = False
    header_row = 1
    width = 0
    heading = ""
    for row_no, values in enumerate(rows, 1):
        cells = [_cell(value) for value in values]
        if header is None:
            if not any(cells):
                continue
            header_row, header, width = row_no, cells, len(cells)
            names = [cell or f"列{index + 1}" for index, cell in enumerate(header)]
            heading = " | ".join(names) + "\n" + " | ".join(["---"] * width)
            continue
        if not any(cells):
            continue
        if len(cells) < width:
            cells.extend([""] * (width - len(cells)))
        elif len(cells) > width:
            # CSV 的异常额外列仍保留，不静默丢弃单元格。
            width = len(cells)
            names = [header[index] if index < len(header) and header[index] else f"列{index + 1}" for index in range(width)]
            heading = " | ".join(names) + "\n" + " | ".join(["---"] * width)
        produced = True
        yield _Block(" | ".join(cells) + "\n", {**base, "row_start": row_no, "row_end": row_no},
                     progress(), "\n".join(part for part in (prefix, heading) if part))
    if header is not None and not produced:
        yield _Block(heading + "\n", {**base, "row_start": header_row, "row_end": header_row}, progress(), prefix)


def _csv_blocks(path):
    with path.open("rb") as stream:
        sample = stream.read(65536)
    encoding = detect_encoding(sample)
    if encoding == "utf-8":
        encoding = "utf-8-sig"
    try:
        dialect = csv.Sniffer().sniff(sample.decode(encoding, errors="replace"), delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    size = max(1, path.stat().st_size)
    with path.open("r", encoding=encoding, errors="replace", newline="") as stream:
        yield from _table_blocks(csv.reader(stream, dialect), {"table_index": 1},
                                 lambda: min(1.0, stream.buffer.tell() / size))


def _excel_blocks(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        total = max(1, len(workbook.worksheets))
        for index, sheet in enumerate(workbook.worksheets):
            yield from _table_blocks(sheet.iter_rows(values_only=True), {"sheet": sheet.title, "table_index": index + 1},
                                     lambda: (index + 1) / total, f"工作表：{sheet.title}")
    finally:
        workbook.close()


def _docx_blocks(path):
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document, sections = Document(str(path)), _Sections()
    children = document.element.body
    total, paragraph_no, table_no = max(1, len(children)), 0, 0
    for index, child in enumerate(children, 1):
        if isinstance(child, CT_P):
            paragraph_no += 1
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            match = re.search(r"(?:Heading|标题)\s*([1-9])", style, re.IGNORECASE)
            sections.update(text, int(match.group(1)) if match else None)
            metadata = sections.metadata() or {"section": f"段落 {paragraph_no}", "section_key": f"p{paragraph_no}"}
            yield _Block(text + "\n", metadata, index / total, sections.prefix())
        elif isinstance(child, CT_Tbl):
            table_no += 1
            table = Table(child, document)
            rows = ([cell.text for cell in row.cells] for row in table.rows)
            yield from _table_blocks(rows, {**sections.metadata(), "table_index": table_no}, lambda: index / total, sections.prefix())


def _blocks(path):
    handlers = {".pdf": _pdf_blocks, ".docx": _docx_blocks, ".txt": _text_blocks, ".md": _text_blocks,
                ".csv": _csv_blocks, ".xlsx": _excel_blocks}
    handler = handlers.get(path.suffix.lower())
    if handler is None:
        raise ValueError("不支持的文档类型")
    yield from handler(path)


def iter_structured_chunks(path: Path, chunk_size: int, chunk_overlap: int):
    if chunk_size <= 0 or not 0 <= chunk_overlap < chunk_size:
        raise ValueError("分块大小或重叠配置不合法")
    buffer, spans, group, prefix, progress = "", [], None, "", 0.0
    produced = False

    def drain(final=False):
        nonlocal buffer, spans, produced
        capacity = max(1, chunk_size - len(prefix) - (1 if prefix else 0))
        while len(buffer) > capacity or (final and buffer):
            cut = min(capacity, len(buffer))
            if len(buffer) > capacity:
                for separator in ("\n\n", "\n", "。", "；", " "):
                    boundary = buffer.rfind(separator, capacity // 2, capacity)
                    if boundary >= 0:
                        cut = boundary + len(separator)
                        break
            active = [metadata for start, end, metadata in spans if start < cut and end > 0]
            body = buffer[:cut]
            if body.strip():
                metadata = {**(active[0] if active else {}), **merge_locations(active)}
                content = (prefix + "\n" if prefix else "") + body
                produced = True
                yield StructuredChunk(content, metadata, progress)
            overlap = 0 if final and cut == len(buffer) else min(chunk_overlap, cut // 2)
            advance = max(1, cut - overlap)
            buffer = buffer[advance:]
            spans = [(max(0, start - advance), end - advance, metadata) for start, end, metadata in spans if end > advance]

    for block in _blocks(Path(path)):
        text = block.text.replace("\x00", "")
        if not text.strip():
            continue
        metadata = dict(block.metadata)
        key = tuple(metadata.get(name) for name in ("section_key", "page_start", "sheet", "table_index"))
        # 表头变化也应另起分块，防止额外列使用上一张表头。
        key += (hashlib.sha256(block.prefix.encode()).hexdigest(),)
        max_prefix = max(0, min(chunk_size // 3, chunk_size - 16))
        if group != key:
            yield from drain(True)
            group = key
            prefix = block.prefix[:max_prefix]
        if len(block.prefix) > max_prefix:
            metadata["prefix_truncated"] = True
        progress = max(progress, min(1.0, block.progress))
        for start in range(0, len(text), max(1, chunk_size * 2)):
            piece = text[start:start + max(1, chunk_size * 2)]
            offset = len(buffer)
            buffer += piece
            spans.append((offset, len(buffer), metadata))
            yield from drain()
    yield from drain(True)
    if not produced:
        raise ValueError("文档内容为空或无法提取文本（扫描件需先做 OCR）")
