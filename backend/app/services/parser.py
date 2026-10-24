"""文档解析：PDF / DOCX / TXT / MD / XLSX / CSV -> 纯文本。

``iter_document_blocks`` 以流式方式产出文本块，避免把 GB 级原文一次性读进内存；
``parse_document`` 保留为兼容入口（内部拼接所有块）。
"""
import codecs
from collections.abc import Iterator
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".xlsx", ".csv"}

# 文本类文件每次读取的字节数
READ_BLOCK = 1024 * 1024
# CSV 分块读取的行数
CSV_ROWS = 50_000
# 编码嗅探的采样字节数
SNIFF_BYTES = 64 * 1024

ENCODINGS = ("utf-8", "utf-8-sig", "gbk", "gb18030")


class UnsupportedFileType(ValueError):
    pass


def _parse_pdf(path: Path) -> str:
    return "".join(text for text, _ in _iter_pdf(path))


def _iter_pdf(path: Path) -> Iterator[tuple[str, float]]:
    import pymupdf

    with pymupdf.open(path) as doc:
        total = doc.page_count or 1
        for page_no, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                yield f"<!-- page {page_no} -->\n{text}\n\n", page_no / total


def _parse_docx(path: Path) -> str:
    return "".join(text for text, _ in _iter_docx(path))


def _iter_docx(path: Path) -> Iterator[tuple[str, float]]:
    from docx import Document as DocxDocument

    # python-docx 需要整体加载压缩包，这是格式限制（见计划"已知限制"）
    doc = DocxDocument(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    total = max(1, len(paragraphs) + len(doc.tables))

    done = 0
    for text in paragraphs:
        done += 1
        yield f"{text}\n", done / total

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        done += 1
        if rows:
            yield "\n".join(rows) + "\n\n", done / total


def _sniff_encoding(path: Path) -> str:
    """按采样字节判断编码，返回第一个能解码的候选。"""
    with path.open("rb") as fh:
        head = fh.read(SNIFF_BYTES)
    for encoding in ENCODINGS:
        try:
            # 采样末尾可能截断多字节字符，去掉末尾不完整序列再判定
            decoder = codecs.getincrementaldecoder(encoding)()
            decoder.decode(head, final=False)
            return encoding
        except UnicodeDecodeError:
            continue
    return "utf-8"


def _parse_text(path: Path) -> str:
    return "".join(text for text, _ in _iter_text(path))


def _iter_text(path: Path) -> Iterator[tuple[str, float]]:
    """按块解码文本文件，内存占用恒定。"""
    size = max(1, path.stat().st_size)
    encoding = _sniff_encoding(path)
    decoder = codecs.getincrementaldecoder(encoding)(errors="replace")

    read = 0
    with path.open("rb") as fh:
        while raw := fh.read(READ_BLOCK):
            read += len(raw)
            text = decoder.decode(raw, final=False)
            if text:
                yield text, read / size
    tail = decoder.decode(b"", final=True)
    if tail:
        yield tail, 1.0


def _parse_tabular(path: Path) -> str:
    return "".join(text for text, _ in _iter_tabular(path))


def _iter_tabular(path: Path) -> Iterator[tuple[str, float]]:
    if path.suffix.lower() == ".csv":
        yield from _iter_csv(path)
    else:
        yield from _iter_excel(path)


def _iter_csv(path: Path) -> Iterator[tuple[str, float]]:
    """CSV 分块读取；进度取自底层字节流位置，大文件进度真实可用。"""
    import io

    import pandas as pd

    size = max(1, path.stat().st_size)
    encoding = _sniff_encoding(path)

    with path.open("rb") as raw:
        stream = io.TextIOWrapper(raw, encoding=encoding, errors="replace", newline="")
        try:
            reader = pd.read_csv(stream, chunksize=CSV_ROWS)
            first = True
            for df in reader:
                text = df.to_markdown(index=False)
                prefix = "" if first else "\n"
                first = False
                yield f"{prefix}{text}\n", min(1.0, raw.tell() / size)
        except pd.errors.EmptyDataError:
            return


def _iter_excel(path: Path) -> Iterator[tuple[str, float]]:
    import pandas as pd

    sheets = pd.read_excel(path, sheet_name=None)
    total = max(1, len(sheets))
    for done, (name, df) in enumerate(sheets.items(), start=1):
        yield f"## 工作表: {name}\n{df.to_markdown(index=False)}\n\n", done / total


ITERATORS = {
    ".pdf": _iter_pdf,
    ".docx": _iter_docx,
    ".txt": _iter_text,
    ".md": _iter_text,
    ".xlsx": _iter_excel,
    ".csv": _iter_csv,
}


def iter_document_blocks(path: Path) -> Iterator[tuple[str, float]]:
    """流式产出 (文本块, 完成度 0~1)。"""
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(f"不支持的文件类型: {ext}")

    produced = False
    for text, fraction in ITERATORS[ext](path):
        cleaned = text.replace("\x00", "")
        if cleaned.strip():
            produced = True
            yield cleaned, fraction
    if not produced:
        raise ValueError("文档内容为空或无法提取文本（扫描件需先做 OCR）")


def parse_document(path: Path) -> str:
    """按扩展名解析文档，返回纯文本（兼容旧调用）。"""
    text = "".join(block for block, _ in iter_document_blocks(path)).strip()
    if not text:
        raise ValueError("文档内容为空或无法提取文本（扫描件需先做 OCR）")
    return text
